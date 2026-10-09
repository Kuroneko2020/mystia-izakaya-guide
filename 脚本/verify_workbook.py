# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
out=[]
out.append('文件: '+str(f))
for ws in wb.worksheets:
    out.append(f'===== {ws.title}: {ws.max_row} 行 x {ws.max_column} 列')
    for i,row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row,4), values_only=True),1):
        vals=[str(v)[:120].replace('\n',' / ') if v is not None else '' for v in row]
        out.append(f'  {i}: '+' | '.join(vals))
(D/'workbook_verify.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

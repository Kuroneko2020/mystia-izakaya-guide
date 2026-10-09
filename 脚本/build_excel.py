# -*- coding: utf-8 -*-
"""东方夜雀食堂客人攻略查询表生成脚本。

数据主来源：夜雀助手（AnYiEE/touhou-mystia-izakaya-assistant）v2.10.11，
其游戏数据更新于 2026-10-04，包含 4.4.0 联动客人秦心。

输出：成品/东方夜雀食堂_客人攻略查询表_v4.4.0_2026-10-04.xlsx
"""
import json
import re
from collections import OrderedDict, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
MID = ROOT / "中间数据"
SRC = ROOT / "原始资料"
ANYI = SRC / "repos" / "touhou-mystia-izakaya-assistant" / "app" / "domain" / "data"
OUT = ROOT / "成品"
OUT.mkdir(parents=True, exist_ok=True)
NL = chr(10)


# ---------------------------------------------------------------- 数据读取与映射
def load_json(name):
    return json.loads((MID / f"{name}.json").read_text(encoding="utf-8"))


def export_object_body(filename, block):
    text = (ANYI / filename).read_text(encoding="utf-8")
    m = re.search(r"export\s+const\s+" + re.escape(block) + r"\s*=\s*\{(.*?)\}\s*as const", text, re.S)
    if not m:
        raise ValueError(f"未找到 {block} in {filename}")
    return m.group(1)


def parse_numeric_tag_map(filename, block):
    body = export_object_body(filename, block)
    return {int(k): v for k, v in re.findall(r"\[?(-?\d+)\]?\s*:\s*'([^']*)'", body)}


FOOD_TAG = parse_numeric_tag_map("tags/tagFacts.ts", "FOOD_TAG_MAP")
BEVERAGE_TAG = parse_numeric_tag_map("tags/tagFacts.ts", "BEVERAGE_TAG_MAP")


def parse_map_facts():
    body = export_object_body("places/placeFacts.ts", "MAP_FACTS")
    out = {}
    for key, dlc, label, tier in re.findall(
        r"(\w+)\s*:\s*\{\s*dlc:\s*([\d.]+)\s*,\s*label:\s*'([^']+)'\s*,\s*unlockTier:\s*(\d+)\s*\}",
        body,
    ):
        out[key] = {"dlc": float(dlc), "label": label, "unlockTier": int(tier)}
    out["Hakugyokurou"] = {"dlc": 0, "label": "白玉楼", "unlockTier": 0}
    return out


MAP = parse_map_facts()


def parse_string_map(filename, block):
    body = export_object_body(filename, block)
    return {k: v for k, v in re.findall(r"(\w+)\s*:\s*'([^']*)'", body)}


MERCHANT = parse_string_map("places/merchantFacts.ts", "MERCHANT_LABEL_MAP")
COLLAB = parse_string_map("labels/collaborationFacts.ts", "COLLABORATION_LABEL_MAP")
ING_TYPE = {-1: "其他", 0: "肉类", 1: "海鲜", 2: "蔬菜"}

COLLECTION_POINTS = {}
for rec in load_json("COLLECTION_POINT_REFRESH_FACTS"):
    for key in (rec.get("label"), *rec.get("labels", [])):
        if key:
            COLLECTION_POINTS[key] = rec.get("displayLabel", key)


def map_text(maps):
    return "、".join(MAP.get(m, {}).get("label", m) for m in maps)


def food_tag_name(tag):
    if tag == -20:
        return "流行喜爱（动态）"
    if tag == -21:
        return "流行厌恶（动态）"
    return FOOD_TAG.get(tag, f"<{tag}>")


def drink_tag_name(tag):
    return BEVERAGE_TAG.get(tag, f"<{tag}>")


def tag_list(tags, func=food_tag_name):
    return "、".join(func(t) for t in tags) if tags else "无"


def fixed_positive(guest):
    return [t for t in guest.get("positiveTags", []) if t != -20]


def fixed_negative(guest):
    return [t for t in guest.get("negativeTags", []) if t != -21]


def dynamic_positive(guest):
    return -20 in guest.get("positiveTags", [])


def dynamic_negative(guest):
    return -21 in guest.get("negativeTags", [])


def format_source(source, guest_names):
    """尽量把 from 字段转成中文来源描述。"""
    if source is None:
        return ""
    if isinstance(source, list):
        return "；".join(x for x in (format_source(s, guest_names) for s in source) if x)
    if not isinstance(source, dict):
        return str(source)
    parts = []
    if source.get("self"):
        parts.append("初始")
    if "levelup" in source:
        lv = source["levelup"].get("level")
        mp = source["levelup"].get("map")
        text = f"餐厅等级Lv{lv}"
        if mp:
            text += f"（{MAP.get(mp, {}).get('label', mp)}）"
        parts.append(text)
    if "bond" in source:
        b = source["bond"]
        gid = b.get("specialGuest")
        gname = guest_names.get(gid, f"ID{gid}")
        parts.append(f"{gname}羁绊Lv{b.get('level')}")
    if "areaTask" in source:
        t = source["areaTask"]
        mp = t.get("map")
        text = f"{MAP.get(mp, {}).get('label', mp)}支线任务"
        if t.get("specialGuest") is not None:
            text += f"（{guest_names.get(t['specialGuest'], t['specialGuest'])}）"
        if t.get("task"):
            text += f"：{t['task']}"
        parts.append(text)
    if "buy" in source:
        names = []

        def walk_buy(x):
            if isinstance(x, dict):
                if "merchant" in x:
                    m = x["merchant"]
                    label = MERCHANT.get(m.get("label"), m.get("label", "?"))
                    if m.get("map"):
                        label += f"（{MAP.get(m['map'], {}).get('label', m['map'])}）"
                    names.append(label)
                elif "specialGuest" in x:
                    label = guest_names.get(x["specialGuest"], f"ID{x['specialGuest']}")
                    if x.get("map"):
                        label += f"（{MAP.get(x['map'], {}).get('label', x['map'])}）"
                    names.append(label)
                elif "label" in x:
                    label = MERCHANT.get(x["label"], x["label"])
                    if x.get("map"):
                        label += f"（{MAP.get(x['map'], {}).get('label', x['map'])}）"
                    names.append(label)
                elif "map" in x:
                    names.append(MAP.get(x["map"], {}).get("label", x["map"]))
            elif isinstance(x, list):
                for y in x:
                    walk_buy(y)

        walk_buy(source["buy"])
        if names:
            parts.append("购买：" + "、".join(dict.fromkeys(names)))
    if "collect" in source:
        names = []

        def walk(x):
            if isinstance(x, dict):
                if x.get("label"):
                    names.append(COLLECTION_POINTS.get(x["label"], x["label"]))
                for key in x.get("labels", []):
                    names.append(COLLECTION_POINTS.get(key, key))
                if x.get("map"):
                    names.append(MAP.get(x["map"], {}).get("label", x["map"]))
            elif isinstance(x, list):
                for y in x:
                    walk(y)

        for x in source["collect"]:
            walk(x)
        if names:
            parts.append("采集：" + "、".join(dict.fromkeys(names)))
    if "fishing" in source or "fishingAdvanced" in source:
        maps = list(source.get("fishing", [])) + list(source.get("fishingAdvanced", []))
        if maps:
            parts.append("钓鱼：" + "、".join(dict.fromkeys(MAP.get(m, {}).get("label", m) for m in maps)))
    if "collaboration" in source:
        c = source["collaboration"]
        parts.append("联动：" + COLLAB.get(c.get("collaborationLabel"), c.get("collaborationLabel", "")))
    if "failedCooking" in source:
        parts.append("料理失败（黑暗物质）")
    if "task" in source:
        tasks = source["task"] if isinstance(source["task"], list) else [source["task"]]
        texts = []
        for t in tasks:
            if isinstance(t, dict):
                if t.get("task"):
                    texts.append(str(t["task"]))
                elif t.get("map"):
                    texts.append("任务：" + MAP.get(t["map"], {}).get("label", t["map"]))
        if texts:
            parts.append("任务：" + "、".join(texts))
    return "；".join([p for p in parts if p]) or json.dumps(source, ensure_ascii=False)


# ---------------------------------------------------------------- 原始数据
special_all = load_json("SPECIAL_GUEST_LIST")
normal_all = load_json("NORMAL_GUEST_LIST")
foods_all = load_json("FOOD_LIST")
ingredients_all = load_json("INGREDIENT_LIST")
beverages_all = load_json("BEVERAGE_LIST")
decorations_all = load_json("DECORATION_LIST")
clothes_all = load_json("CLOTHES_LIST")

TAG_COVERS = [(-3, -2), (-1, 28), (0, 2), (6, 7), (9, 8), (22, 21)]


def process_food(d):
    """按夜雀助手 FoodCatalog 口径补价格动态标签并处理标签覆盖。"""
    pos = set(d.get("positiveTags", []))
    if d.get("id") != -1:
        if d.get("price", 0) > 60:
            pos.add(-3)
        elif d.get("price", 0) < 20:
            pos.add(-2)
    for coverer, covered in TAG_COVERS:
        if coverer in pos:
            pos.discard(covered)
    d = dict(d)
    d["positiveTags"] = sorted(pos)
    return d


special = [r for r in special_all if r.get("dlc") != 9]
normal = [r for r in normal_all if r.get("dlc") != 9]
foods = [process_food(r) for r in foods_all if r.get("dlc") != 9 and r.get("id") != -1]
ingredients = [r for r in ingredients_all if r.get("dlc") != 9]
beverages = [r for r in beverages_all if r.get("dlc") != 9]
decorations = [r for r in decorations_all if r.get("dlc") != 9]
clothes = [r for r in clothes_all if r.get("dlc") != 9]

guest_names = {g["id"]: g["name"] for g in special}
ing_map = {r["id"]: r for r in ingredients}
ing_tags = {r["id"]: set(r.get("tags", [])) for r in ingredients}

SPECIAL_TYPE_NAMES = {
    "雾雨魔理沙",
    "森近霖之助",
    "魂魄妖梦",
    "西行寺幽幽子",
    "蕾米莉亚",
    "饕餮尤魔",
    "古明地恋",
    "二岩猯藏",
    "八云紫",
}


def dlc_label(dlc):
    if dlc == 0:
        return "本篇"
    if dlc == 2.5:
        return "DLC2.5"
    if dlc == 9:
        return "MetaMystia示例包（非官方）"
    return f"DLC{dlc}"


def guest_type(g):
    if g.get("collaboration"):
        return "联动角色"
    if g["name"] in SPECIAL_TYPE_NAMES:
        return "特殊顾客"
    if g.get("dlc", 0) == 0:
        return "本篇稀客"
    return f"{dlc_label(g.get('dlc'))}稀客"


def guest_bond_text(g):
    gid = g["id"]
    items = defaultdict(list)
    for f in foods:
        b = f.get("from", {}).get("bond") if isinstance(f.get("from"), dict) else None
        if b and b.get("specialGuest") == gid:
            items[b.get("level")].append(f["name"] + "（菜谱）")
    for d in decorations:
        b = d.get("from", {}).get("bond") if isinstance(d.get("from"), dict) else None
        if b and b.get("specialGuest") == gid:
            items[b.get("level")].append(d["name"] + "（摆件）")
    for c in clothes:
        f = c.get("from")
        entries = f if isinstance(f, list) else [f]
        for entry in entries:
            if isinstance(entry, dict) and entry.get("bond", {}).get("specialGuest") == gid:
                items[entry["bond"].get("level")].append(c["name"] + "（衣服）")
    if not items:
        return ""
    return "；".join(f"Lv{lv}：" + "、".join(vals) for lv, vals in sorted(items.items(), key=lambda x: (x[0] is None, x[0])))


def spell_text(g):
    cards = g.get("spellCards") or {}
    pos = "；".join(f"{c['name']}：{c['description']}" for c in cards.get("positive", []))
    neg = "；".join(f"{c['name']}：{c['description']}" for c in cards.get("negative", []))
    return pos or "无", neg or "无"


def dish_recipe_text(d):
    recipes = d.get("recipes") or []
    if not recipes:
        return "—", "—"
    cookers = []
    mats = []
    for r in recipes:
        cooker = {1: "煮锅", 2: "烧烤架", 3: "油锅", 4: "蒸锅", 5: "料理台"}.get(r.get("cookerType"), f"厨具{r.get('cookerType')}")
        cookers.append(cooker)
        names = [ing_map[i]["name"] for i in r.get("ingredients", []) if i in ing_map]
        mats.append("＋".join(names) if names else "无")
    return "、".join(dict.fromkeys(cookers)), "；".join(dict.fromkeys(mats))


def dish_ingredient_ids(d):
    ids = []
    for r in d.get("recipes") or []:
        ids += r.get("ingredients", [])
    return list(dict.fromkeys(ids))


def first_recipe_ids(d):
    recipes = d.get("recipes") or []
    return recipes[0].get("ingredients", []) if recipes else []


def dish_score(dish, guest):
    pos = set(dish.get("positiveTags", []))
    likes = set(fixed_positive(guest))
    dislikes = set(fixed_negative(guest))
    mapping = {int(k) for k in (guest.get("positiveTagMapping") or {}).keys()}
    like_hits = len(pos & likes)
    dislike_hits = len(pos & dislikes)
    mapping_hits = len(pos & mapping)
    dynamic_hits = 0
    if dynamic_positive(guest) and -20 in pos:
        dynamic_hits += 1
    if dynamic_negative(guest) and -21 in pos:
        dynamic_hits -= 1
    return {
        "like_hits": like_hits,
        "dislike_hits": dislike_hits,
        "mapping_hits": mapping_hits,
        "dynamic_hits": dynamic_hits,
        "game_score": like_hits - dislike_hits,
        "pos": pos,
        "likes": likes,
        "dislikes": dislikes,
        "mapping": mapping,
    }


def early_ok(d, guest, relax=False):
    if d.get("dlc") != 0:
        return False
    src = d.get("from")
    if isinstance(src, dict) and "collaboration" in src:
        label = src["collaboration"].get("collaborationLabel")
        if guest.get("collaboration") != label:
            return False
    if not relax:
        if d.get("level", 0) > 3:
            return False
        if d.get("price", 0) > min(guest["price"][1], 120):
            return False
        if isinstance(src, dict):
            if "bond" in src and src["bond"].get("level", 0) > 3:
                return False
            if "levelup" in src and src["levelup"].get("level", 0) > 3:
                return False
        for iid in first_recipe_ids(d):
            ing = ing_map.get(iid)
            if not ing or ing.get("dlc") != 0 or ing.get("level", 0) > 3:
                return False
    else:
        if d.get("price", 0) > guest["price"][1]:
            return False
    return True


def late_ok(d, guest, min_like=3, relax=False):
    if d.get("price", 0) > guest["price"][1]:
        return False
    if d.get("dlc", 0) > 5:
        return False
    s = dish_score(d, guest)
    if s["dislike_hits"] > 0:
        return False
    if not relax and (s["like_hits"] < min_like and s["mapping_hits"] < min_like):
        return False
    return True


def select_dishes(guest):
    scored = []
    for d in foods:
        s = dish_score(d, guest)
        scored.append((d, s))
    early = [(d, s) for d, s in scored if early_ok(d, guest, False)]
    early = [x for x in early if x[1]["dislike_hits"] == 0 and x[1]["like_hits"] >= 2]
    early.sort(key=lambda x: (-x[1]["game_score"], -x[1]["mapping_hits"], x[0]["price"], x[0]["level"]))
    if len(early) < 2:
        early = [(d, s) for d, s in scored if early_ok(d, guest, True)]
        early = [x for x in early if x[1]["dislike_hits"] == 0 and x[1]["like_hits"] >= 1]
        early.sort(key=lambda x: (-x[1]["game_score"], -x[1]["mapping_hits"], x[0]["price"], x[0]["level"]))
    early = early[:3]

    late = [(d, s) for d, s in scored if late_ok(d, guest, 3, False)]
    if len(late) < 2:
        late = [(d, s) for d, s in scored if late_ok(d, guest, 2, True)]
    if len(late) < 2:
        late = [(d, s) for d, s in scored if late_ok(d, guest, 1, True)]
    late.sort(key=lambda x: (-x[1]["game_score"], -x[1]["mapping_hits"], -x[0]["price"], -x[0]["level"]))
    early_ids = {d["id"] for d, _ in early}
    late_unique = [x for x in late if x[0]["id"] not in early_ids]
    if len(late_unique) >= 2:
        late = late_unique
    late = late[:3]
    return early, late


def addon_candidates(guest, dish=None):
    likes = set(fixed_positive(guest))
    dislikes = set(fixed_negative(guest))
    used = set(first_recipe_ids(dish)) if dish else set()
    conflicts = set(dish.get("negativeTags", [])) if dish else set()
    out = []
    for ing in ingredients:
        if ing.get("id") == -1 or ing.get("price", 0) > 10000:
            continue
        tags = set(ing.get("tags", []))
        if tags & dislikes or tags & conflicts:
            continue
        if ing["id"] in used:
            continue
        hit = tags & likes
        if not hit:
            continue
        out.append((len(hit), len(tags), ing, hit))
    out.sort(key=lambda x: (-x[0], -x[1], x[2].get("dlc", 0), x[2].get("level", 0), x[2].get("price", 0)))
    return out


def addon_text_for_guest(guest):
    cands = addon_candidates(guest)
    if not cands:
        return "暂无可补喜好标签的加料；优先选已满足喜好标签较多的菜。"
    parts = [f"{ing['name']}（{tag_list(sorted(hit))}）" for _, _, ing, hit in cands[:6]]
    return "通用加料：" + "、".join(parts) + "。注意：不要加厌恶标签对应食材。"


def addon_text_for_dish(guest, dish):
    s = dish_score(dish, guest)
    missing_all = s["likes"] - s["pos"]
    dynamic_missing = [t for t in missing_all if t in (-20, -21, -2, -3)]
    missing = sorted(t for t in missing_all if t not in (-20, -21, -2, -3))
    if not missing:
        if dynamic_missing:
            return "固定喜好已覆盖；" + tag_list(dynamic_missing) + "为动态标签，不靠加料（看价格/剪报）。"
        return "菜品已覆盖全部固定喜好标签，通常无需加料。"
    cands = []
    for n, _n, ing, hit in addon_candidates(guest, dish):
        real_hit = hit & set(missing)
        if real_hit:
            cands.append((len(real_hit), ing, real_hit))
    if not cands:
        text = "缺：" + tag_list(missing) + "；该菜没有不冲突的加料位，可换其他推荐菜或直接用高命中菜。"
    else:
        cands.sort(key=lambda x: (-x[0], x[1].get("price", 0)))
        sample = "；".join(f"+{ing['name']}（补{tag_list(sorted(hit))}）" for _, ing, hit in cands[:3])
        text = "缺：" + tag_list(missing) + "；可加：" + sample
    if dynamic_missing:
        text += "（" + tag_list(dynamic_missing) + "为动态标签，不靠加料）"
    return text


def recommend_drinks(guest):
    scored = []
    for b in beverages:
        if b.get("price", 0) > guest["price"][1]:
            continue
        hits = set(b.get("tags", [])) & set(guest.get("beverageTags", []))
        if hits:
            scored.append((len(hits), b, hits))
    scored.sort(key=lambda x: (-x[0], x[1].get("price", 0), x[1].get("level", 0)))
    return scored[:3]


def dish_brief(d, s):
    dlc = dlc_label(d.get("dlc", 0))
    return f"{d['name']}（{d['price']}｜{dlc}｜{tag_list(sorted(s['pos']))}）"


# ---------------------------------------------------------------- 解析 B 站攻略
def normalize_name(name):
    name = name.strip()
    name = re.sub(r"[（(].*", "", name).strip()
    return name


def parse_bili_2024():
    text = (SRC / "bili_opus_665676349889839109.txt").read_text(encoding="utf-8")
    lines = [x.strip() for x in text.splitlines()]
    aliases = {"二岩瑞藏": "二岩猯藏", "路易兹": "露易兹"}
    known = {g["name"] for g in special} | set(aliases)
    sections = {
        "妖怪兽道", "人间之里", "博丽神社", "红魔馆", "迷途竹林（包含DLC4角色）",
        "魔法森林DLC1", "妖怪之山DLC1", "旧地狱DLC2", "地灵殿DLC2", "特殊顾客",
        "命莲寺DLC3", "神灵庙DLC3", "辉针城DLC4", "太阳花田DLC4", "魔界 DLC5", "月之都 DLC5",
    }
    out = []
    cur = None
    field = None
    region = ""
    started = False
    for raw in lines:
        if raw == "妖怪兽道" and not started:
            started = True
        if not started:
            continue
        if raw in sections:
            region = raw
            cur = None
            field = None
            continue
        nm = normalize_name(raw)
        if nm in known or nm in aliases:
            cur = OrderedDict([
                ("地区", region), ("客人", aliases.get(nm, nm)),
                ("持有金(攻略)", ""), ("常出没(攻略)", ""),
                ("喜好(攻略)", ""), ("厌恶(攻略)", ""), ("酒水(攻略)", ""),
                ("菜品推荐(攻略原文)", ""), ("酒水推荐(攻略原文)", ""),
                ("羁绊解锁(攻略原文)", ""), ("备注/版本变化", ""),
            ])
            out.append(cur)
            field = None
            continue
        if cur is None:
            continue
        m = re.match(r"^(持有金|常出没|喜|厌|酒水)\s*[:：]\s*(.*)$", raw)
        if m:
            field = m.group(1)
            key = {"喜": "喜好(攻略)", "厌": "厌恶(攻略)"}.get(field, field + "(攻略)")
            cur[key] = m.group(2)
            continue
        if raw in ("【菜品】", "【酒水】", "【羁绊等级解锁】"):
            field = raw
            continue
        if raw.startswith("DLC5") or "瑞灵附身版" in raw:
            cur["备注/版本变化"] = (cur["备注/版本变化"] + "；" + raw).strip("；")
            field = "备注"
            continue
        if field in ("持有金", "常出没", "喜", "厌", "酒水"):
            key = {"喜": "喜好(攻略)", "厌": "厌恶(攻略)"}.get(field, field + "(攻略)")
            cur[key] = (cur.get(key, "") + raw).strip()
        elif field == "【菜品】":
            cur["菜品推荐(攻略原文)"] = (cur["菜品推荐(攻略原文)"] + "；" + raw).strip("；")
        elif field == "【酒水】":
            cur["酒水推荐(攻略原文)"] = (cur["酒水推荐(攻略原文)"] + "；" + raw).strip("；")
        elif field == "【羁绊等级解锁】":
            cur["羁绊解锁(攻略原文)"] = (cur["羁绊解锁(攻略原文)"] + "；" + raw).strip("；")
        elif field == "备注":
            cur["备注/版本变化"] = (cur["备注/版本变化"] + "；" + raw).strip("；")
    return out


def parse_bili_2021():
    text = (SRC / "bili_opus_579630987930616350.txt").read_text(encoding="utf-8")
    lines = [x.strip() for x in text.splitlines()]
    aliases = {"魔理沙": "雾雨魔理沙", "霖之助": "森近霖之助", "妖梦": "魂魄妖梦",
               "幽幽子": "西行寺幽幽子", "三妖精": "蹦蹦跳跳的三妖精"}
    known = {g["name"] for g in special} | set(aliases)
    out = []
    cur = None
    field = None
    separators = [i for i, line in enumerate(lines) if "————" in line]
    start_idx = separators[0] if separators else 0
    for raw in lines[start_idx + 1:]:
        nm = normalize_name(raw)
        if nm in known or nm in aliases:
            cur = OrderedDict([
                ("客人", aliases.get(nm, nm)), ("持有金(攻略)", ""), ("符卡(攻略)", ""),
                ("喜好(攻略)", ""), ("厌恶(攻略)", ""), ("酒水偏好(攻略)", ""),
                ("菜品推荐(攻略原文)", ""), ("加料推荐(攻略原文)", ""),
                ("酒水推荐(攻略原文)", ""), ("羁绊解锁(攻略原文)", ""),
            ])
            out.append(cur)
            field = None
            continue
        if cur is None:
            continue
        m = re.match(r"^[（(]持有金\s*[:：]?\s*([^）)]*)[）)]$", raw)
        if m:
            cur["持有金(攻略)"] = m.group(1)
            continue
        m = re.match(r"^[（(]符卡效果\s*[:：]\s*(.*)$", raw)
        if m:
            cur["符卡(攻略)"] = m.group(1).rstrip("）)")
            continue
        m = re.match(r"^(喜|厌)\s*[:：]\s*(.*)$", raw)
        if m:
            field = m.group(1)
            key = {"喜": "喜好(攻略)", "厌": "厌恶(攻略)"}[field]
            cur[key] = m.group(2)
            continue
        m = re.match(r"^酒水偏好\s*[:：]\s*(.*)$", raw)
        if m:
            field = "酒水偏好"
            cur["酒水偏好(攻略)"] = m.group(1)
            continue
        if raw in ("【菜品】", "【加料】", "【酒水】", "【羁绊等级解锁】"):
            field = raw
            continue
        if field in ("喜", "厌", "酒水偏好"):
            key = {"喜": "喜好(攻略)", "厌": "厌恶(攻略)"}.get(field, field + "(攻略)")
            cur[key] = (cur.get(key, "") + raw).strip()
        elif field == "【菜品】":
            cur["菜品推荐(攻略原文)"] = (cur["菜品推荐(攻略原文)"] + "；" + raw).strip("；")
        elif field == "【加料】":
            cur["加料推荐(攻略原文)"] = (cur["加料推荐(攻略原文)"] + "；" + raw).strip("；")
        elif field == "【酒水】":
            cur["酒水推荐(攻略原文)"] = (cur["酒水推荐(攻略原文)"] + "；" + raw).strip("；")
        elif field == "【羁绊等级解锁】":
            cur["羁绊解锁(攻略原文)"] = (cur["羁绊解锁(攻略原文)"] + "；" + raw).strip("；")
    return out


# ---------------------------------------------------------------- Excel 输出
def clean_cell(value):
    if value is None:
        return ""
    s = str(value).replace("{{br}}", NL)
    s = "".join(ch for ch in s if ch in (NL, "\t") or ord(ch) >= 32)
    return s[:30000]


def add_sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title=title)
    ws.append([clean_cell(h) for h in headers])
    for row in rows:
        ws.append([clean_cell(x) for x in row])
    header_fill = PatternFill("solid", fgColor="305496")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    bili2024 = parse_bili_2024()
    bili2024_map = {r["客人"]: r for r in bili2024}
    bili2021 = parse_bili_2021()
    bili2021_map = {r["客人"]: r for r in bili2021}

    guest_rows = []
    detail_rows = []
    normal_rows = []
    for g in special:
        early, late = select_dishes(g)
        reward, punish = spell_text(g)
        bond = guest_bond_text(g) or ("无（联动角色）" if g.get("collaboration") else "无")
        early_text = NL.join(dish_brief(d, s) for d, s in early) or "无"
        late_text = NL.join(dish_brief(d, s) for d, s in late) or "无"
        addon_text = addon_text_for_guest(g)
        examples = []
        if early:
            examples.append("前期示例：" + addon_text_for_dish(g, early[0][0]))
        if late:
            examples.append("后期示例：" + addon_text_for_dish(g, late[0][0]))
        if examples:
            addon_text += NL + NL.join(examples)
        drinks = recommend_drinks(g)
        drink_text = "；".join(f"{b['name']}（{b['price']}｜{tag_list(sorted(hit), drink_tag_name)}）" for _, b, hit in drinks) or "无"
        guide = bili2024_map.get(g["name"], {})
        guide_text = guide.get("菜品推荐(攻略原文)", "") or "无（该攻略早于4.4.0）"
        version_label = dlc_label(g.get("dlc", 0))
        if g.get("collaboration"):
            version_label = ("4.4.0 " if g["name"] == "秦心" else "") + "联动"
        known_notes = {
            "火焰猫燐": "持有金以夜雀助手500-700为准（THBWiki为500-600）",
            "茨木华扇": "出没地点以夜雀助手为准（B站2024攻略写魔法森林；THBWiki与助手为旧地狱）",
            "西行寺幽幽子": "持有金以夜雀助手1500-2000为准（B站2024攻略为1300-1800）",
            "秦心": "4.4.0（2026-04-17）新增《东方华心传》联动稀客；B站首篇攻略称出没博丽神社/妖怪山、预算约800，本表采用夜雀助手最新数据（5处出没、514-805）",
            "二岩猯藏": "B站旧攻略写作“二岩瑞藏”，本表按正式译名",
            "露易兹": "B站攻略写作“路易兹”，夜雀助手为“露易兹”",
        }
        note_bits = ["数据：夜雀助手 v2.10.11（2026-10-04）"]
        if g.get("collaboration"):
            note_bits.append("联动：" + COLLAB.get(g["collaboration"], g["collaboration"]))
        if g["name"] == "秦心":
            note_bits.append("4.4.0新增联动稀客")
        if g["name"] in known_notes:
            note_bits.append(known_notes[g["name"]])
        if guide:
            note_bits.append("B站2024攻略持有金：" + guide.get("持有金(攻略)", ""))
        guest_rows.append([
            g["name"], guest_type(g), version_label, map_text(g["maps"]),
            f"{g['price'][0]}-{g['price'][1]}" if g["price"][0] != g["price"][1] else f"{g['price'][0]}（固定）",
            tag_list(g["positiveTags"]), tag_list(g["negativeTags"]), tag_list(g["beverageTags"], drink_tag_name),
            early_text, late_text, addon_text, drink_text,
            "；".join(f"{FOOD_TAG.get(int(k), k)}：{v}" for k, v in (g.get("positiveTagMapping") or {}).items()) or "无",
            reward, punish, bond, guide_text, "；".join(note_bits),
        ])
        for stage, items in (("前期", early), ("后期", late)):
            for d, s in items:
                detail_rows.append([
                    g["name"], guest_type(g), stage, d["name"], d["price"], dlc_label(d.get("dlc", 0)),
                    dish_recipe_text(d)[0], dish_recipe_text(d)[1], tag_list(sorted(s["pos"])),
                    tag_list(sorted(s["pos"] & s["likes"])), tag_list(sorted(s["pos"] & s["mapping"])),
                    s["game_score"], format_source(d.get("from"), guest_names),
                    addon_text_for_dish(g, d),
                    "含流行标签，需看当日剪报" if (-20 in s["pos"] or -21 in s["pos"]) else "",
                ])
    for g in normal:
        early, late = [], []
        for d in foods:
            s = dish_score(d, g)
            if s["like_hits"]:
                if d.get("dlc") == 0 and d.get("level", 0) <= 3 and d["price"] <= 80 and s["dislike_hits"] == 0:
                    early.append((d, s))
                if d["price"] <= 200 and s["dislike_hits"] == 0:
                    late.append((d, s))
        early.sort(key=lambda x: (x[1]["like_hits"], x[1]["mapping_hits"], -x[0]["price"]), reverse=True)
        late.sort(key=lambda x: (x[1]["like_hits"], x[1]["mapping_hits"], x[0]["price"]), reverse=True)
        early_text = NL.join(f"{d['name']}（{d['price']}）" for d, _ in early[:3]) or "无"
        late_text = NL.join(f"{d['name']}（{d['price']}）" for d, _ in late[:3]) or "无"
        normal_rows.append([
            g["name"], dlc_label(g.get("dlc", 0)), map_text(g["maps"]), tag_list(g["positiveTags"]),
            tag_list(g.get("beverageTags", []), drink_tag_name), early_text, late_text,
        ])

    wb = Workbook()
    wb.remove(wb.active)

    # 目录
    toc = [
        ("客人总表", "按客人查询：出没/金钱 → 喜好/厌恶 → 前期/后期菜谱 → 加料/酒水 → 符卡/羁绊。"),
        ("推荐菜谱明细", "每条推荐菜谱一行，含评分、标签命中、材料、厨具、加料建议、解锁来源。"),
        ("菜谱速查", "全部官方菜谱（DLC0-5）的标签、价格、厨具、材料和解锁来源。"),
        ("食材速查", "全部官方食材的标签、价格、等级和获取方式。"),
        ("酒水速查", "全部官方酒水的标签、价格和获取方式。"),
        ("普通客人", "47 种普通客人的出没、喜好标签和推荐菜。"),
        ("4.4.0新增内容", "v4.4.0（2026-04-17）新增的秦心、四个情绪面具、妖精服、菜单预设和剧情优化。"),
        ("B站2024攻略对照", "B站稀客图鉴（v4.0.3a，更新至DLC5）原文结构化对照；不含4.4.0秦心。"),
        ("B站2021攻略对照", "B站早期攻略（v1.5.x）的喜好与加料推荐，用于前期基础客人对照。"),
        ("版本与数据来源", "版本、来源、统计口径、核验差异和免责声明。"),
    ]
    ws = wb.create_sheet("目录")
    ws.append(["工作表", "内容"])
    for row in toc:
        ws.append(row)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    add_sheet(
        wb, "客人总表",
        ["客人名字", "类型", "所属版本", "出没地点", "持有金", "喜好料理标签", "厌恶料理标签",
         "喜好酒水标签", "前期推荐菜谱", "后期推荐菜谱", "加料推荐", "推荐酒水",
         "拿手菜/额外喜好标签", "奖励符卡", "惩罚符卡", "羁绊等级解锁",
         "B站2024攻略菜品（对照，原文）", "备注/来源"],
        guest_rows,
        widths=[12, 12, 10, 26, 12, 26, 22, 20, 40, 40, 40, 32, 30, 45, 45, 36, 45, 30],
    )
    add_sheet(
        wb, "推荐菜谱明细",
        ["客人名字", "类型", "阶段", "菜品", "价格", "所属版本", "厨具", "主料", "菜品标签",
         "命中喜好", "命中拿手标签", "游戏积分", "解锁来源", "加料建议", "备注"],
        detail_rows,
        widths=[12, 12, 8, 16, 8, 10, 10, 26, 30, 22, 20, 10, 28, 42, 24],
    )
    dish_rows = []
    for d in foods:
        cooker, mats = dish_recipe_text(d)
        desc = d.get("description", "")
        if d.get("id") == -1:
            desc = "【黑暗物质：料理失败产物，不可作为正常推荐】" + desc
        dish_rows.append([
            d.get("id"), d["name"], dlc_label(d.get("dlc", 0)), d.get("level", ""), d.get("price", ""),
            cooker, mats, tag_list(d.get("positiveTags", [])), tag_list(d.get("negativeTags", [])),
            format_source(d.get("from"), guest_names), desc,
        ])
    add_sheet(
        wb, "菜谱速查",
        ["菜品ID", "菜品", "所属版本", "等级", "价格", "厨具", "主料", "菜品标签",
         "冲突标签（加料会变黑暗物质）", "解锁来源", "描述"],
        dish_rows,
        widths=[8, 18, 10, 8, 8, 10, 30, 34, 24, 28, 50],
    )
    ing_rows = []
    for r in ingredients:
        desc = r.get("description", "")
        if r.get("id") == -1:
            desc = "【特殊彩蛋物品，价格极高，不建议作为常规加料】" + desc
        ing_rows.append([
            r.get("id"), r["name"], ING_TYPE.get(r.get("type"), r.get("type", "")), tag_list(r.get("tags", [])), dlc_label(r.get("dlc", 0)),
            r.get("level", ""), r.get("price", ""), format_source(r.get("from"), guest_names), desc,
        ])
    add_sheet(
        wb, "食材速查",
        ["食材ID", "食材", "类型", "标签", "所属版本", "等级", "价格", "获取方式", "描述"],
        ing_rows,
        widths=[8, 14, 8, 24, 10, 8, 8, 50, 46],
    )
    drink_rows = []
    for r in beverages:
        drink_rows.append([
            r.get("id"), r["name"], dlc_label(r.get("dlc", 0)), r.get("level", ""), r.get("price", ""),
            tag_list(r.get("tags", []), drink_tag_name), format_source(r.get("from"), guest_names), r.get("description", ""),
        ])
    add_sheet(
        wb, "酒水速查",
        ["酒水ID", "酒水", "所属版本", "等级", "价格", "标签", "获取方式", "描述"],
        drink_rows,
        widths=[8, 18, 10, 8, 8, 28, 46, 46],
    )
    add_sheet(
        wb, "普通客人",
        ["普通客人", "所属版本", "出没地点", "喜好料理标签", "喜好酒水标签", "前期推荐菜（≤80元）", "后期推荐菜（≤200元）"],
        normal_rows,
        widths=[12, 10, 30, 36, 24, 26, 26],
    )
    add_sheet(
        wb, "4.4.0新增内容",
        ["类别", "名称", "效果/说明", "获取/备注"],
        [
            ("联动稀客", "秦心", "出没：人间之里、博丽神社、妖怪之山、命莲寺、神灵庙；持有金514-805；喜好：传说、和风、招牌、适合拍照、凉爽、梦幻；厌恶：重油、咸、灼热；符卡：华心「喜怒哀乐 变幻」及喜/怒/哀/乐四种强化。", "《东方华心传》联动；在门口终端领取联动内容。"),
            ("情绪摆件", "喜之面", "客人预算提高20%；被秦心符卡强化后提高到200%。", "《东方华心传》联动；情绪面具摆件只能选择一个摆放。"),
            ("情绪摆件", "怒之面", "料理时间减少10%；被秦心符卡强化后提高到75%。", "《东方华心传》联动；情绪面具摆件只能选择一个摆放。"),
            ("情绪摆件", "哀之面", "客人续单概率提高5%；被秦心符卡强化后提高到50%。", "《东方华心传》联动；情绪面具摆件只能选择一个摆放。"),
            ("情绪摆件", "乐之面", "稀有客人首次评价至少为普通；被秦心符卡强化后至少为完美。", "《东方华心传》联动；情绪面具摆件只能选择一个摆放。"),
            ("衣服", "妖精服", "《东方妖精武踏会》联动服装，更换时店铺夜间装潢变为妖精风格。", "在门口终端领取联动内容。"),
            ("系统", "菜单预设", "新增菜单预设与改名功能，便于快速切换菜单。", "v4.4.0新增。"),
            ("剧情", "DLC1/DLC2辅助任务", "缩短DLC1主线事件等待时间，补充失败演出；增加DLC2旧地狱开店辅助任务等。", "v4.4.0优化。"),
        ],
        widths=[12, 14, 80, 36],
    )
    bili2024_headers = ["地区", "客人", "持有金(攻略)", "常出没(攻略)", "喜好(攻略)", "厌恶(攻略)", "酒水(攻略)",
                        "菜品推荐(攻略原文)", "酒水推荐(攻略原文)", "羁绊解锁(攻略原文)", "备注/版本变化"]
    add_sheet(
        wb, "B站2024攻略对照", bili2024_headers,
        [[r.get(h, "") for h in bili2024_headers] for r in bili2024],
        widths=[14, 12, 12, 24, 26, 22, 22, 60, 50, 45, 24],
    )
    bili2021_headers = ["客人", "持有金(攻略)", "符卡(攻略)", "喜好(攻略)", "厌恶(攻略)", "酒水偏好(攻略)",
                        "菜品推荐(攻略原文)", "加料推荐(攻略原文)", "酒水推荐(攻略原文)", "羁绊解锁(攻略原文)"]
    add_sheet(
        wb, "B站2021攻略对照", bili2021_headers,
        [[r.get(h, "") for h in bili2021_headers] for r in bili2021],
        widths=[12, 12, 34, 26, 22, 22, 50, 40, 40, 40],
    )

    version_rows = [
        ("表格生成日期", "2026-10-04"),
        ("主数据版本", "夜雀助手（AnYiEE/touhou-mystia-izakaya-assistant）v2.10.11；仓库最近提交 2026-10-04"),
        ("游戏内容版本", "官方最新内容公告为 v4.4.0（2026-04-17），新增《东方华心传》联动稀客秦心、四个情绪面具摆件，以及《东方妖精武踏会》妖精服。SteamDB 另记录 2026-05-09 有一次更新（https://steamdb.info/patchnotes/23158340/），截至 2026-10-04 未检索到新增官方客人/DLC内容的公告。"),
        ("DLC范围", "官方 DLC0（本篇）至 DLC5（魔界/月之都），共 59 名官方特殊/稀客（含特殊顾客与联动）、47 种普通客人、163 道官方菜谱、46 款官方酒水、62 种官方食材（已剔除 MetaMystia 示例包的 DLC9 模组数据；另有 1 条“黑暗物质”失败料理未列入菜谱表）。"),
        ("主数据来源1", "夜雀助手 GitHub：https://github.com/AnYiEE/touhou-mystia-izakaya-assistant （数据文件 app/domain/data）"),
        ("主数据来源2", "夜雀助手网站：https://izakaya.cc"),
        ("数据许可与致谢", "主数据来自夜雀助手（AnYiEE/touhou-mystia-izakaya-assistant，AGPL-3.0-only）；本表为玩家查询用途的数据整理与再加工，非官方资料。游戏内容版权归 二色幽紫蝶 / Re零同人社 等相关权利方所有。"),
        ("版本公告", "Steam 公告 Important Update: Version 4.4.0（2026-04-17）：https://store.steampowered.com/news/app/1584090"),
        ("中文更新说明", "17173 转载：https://news.17173.com/content/04182026/062526887.shtml"),
        ("交叉核对1", "THBWiki《攻略/客人》：https://thbwiki.cc/东方夜雀食堂_-_Touhou_Mystia%27s_Izakaya/攻略/客人 （最后编辑 2025-11-08，用于核对本体至DLC5的稀有/普通/特殊顾客）"),
        ("交叉核对2", "B站稀客图鉴及推荐菜单（v4.0.3a，更新至DLC5）：https://m.bilibili.com/opus/665676349889839109 （2024-03-01）"),
        ("交叉核对3", "B站早期稀客图鉴（v1.5.x，含加料专列）：https://m.bilibili.com/opus/579630987930616350 （2021-12-11）"),
        ("交叉核对4", "Gitee 夜雀助手（2024-06-01 数据）：https://gitee.com/troublemove/touhou-izakaya-helper （用于DLC5前数据核对）"),
        ("秦心早期信息", "B站专栏 cv47914353《全网首发！东方夜雀食堂4月17日更新特殊顾客秦心Tag》（2026-04-18）：https://www.bilibili.com/read/cv47914353 。该文称秦心出没博丽神社/妖怪山、预算约800；本表采用夜雀助手最新数据（出没人间之里、博丽神社、妖怪之山、命莲寺、神灵庙；持有金514-805）。"),
        ("统计口径-前期/后期", "前期：本体DLC0、菜谱等级≤3、单价≤120且材料本体早期可获取；优先命中喜好标签≥2、无厌恶标签。后期：单价不超过客人持有金上限、命中喜好/拿手标签≥3，优先高收益菜；若不足则自动放宽并仍在表中标注。此分级为查询用规则，不是游戏官方分级。"),
        ("统计口径-积分", "按游戏规则：菜品每命中一个喜好标签 +1，命中一个厌恶标签 -1；拿手标签单独列出，用于提高开奖励符卡概率。流行喜爱/厌恶随每日剪报变化，本表静态列出但评分不计入流行标签。"),
        ("统计口径-菜品标签", "菜品标签按夜雀助手 FoodCatalog 口径处理：单价>60补“昂贵”，单价<20补“实惠”；肉覆盖素、重油覆盖清淡、饱腹覆盖下酒、灼热覆盖凉爽、大份覆盖小巧。流行喜爱/厌恶为每日动态标签，本表不把流行标签计入固定评分。"),
        ("统计口径-加料", "加料候选=食材标签命中客人喜好、且不与客人厌恶标签和该菜冲突标签冲突；需避开会覆盖原标签或产生黑暗物质的组合。实际加料还受菜品已有标签和料理小游戏影响。"),
        ("统计口径-普通客人", "普通客人只有喜好标签、没有厌恶标签与固定持有金；推荐菜按喜好命中数排序，价格上限为查询用的经验值（前期≤80、后期≤200）。"),
        ("统计口径-持有金", "持有金为客人当晚预算范围，推荐菜按“单道菜价格不超过该上限”筛选；部分客人预算浮动（如秦心514-805）。实际以游戏内显示为准。"),
        ("统计口径-随机点单", "稀客点单包含随机点单标签；本表推荐菜只保证静态喜好/厌恶匹配，实际开奖励符卡还需满足点单要求，夜雀系列厨具的完美料理可无视点单标签。"),
        ("加料提示", "加料候选按命中标签数、标签总数、DLC、等级和价格排序，优先本体低等级食材；若只列出高等级/DLC食材，说明该标签在该阶段较稀缺，请结合“食材速查”确认获取方式。"),
        ("差异1-火焰猫燐持有金", "THBWiki记500-600；B站2024攻略与夜雀助手均为500-700。本表采用夜雀助手最新数据500-700。"),
        ("差异2-茨木华扇出没", "B站2024攻略写人间之里/魔法森林/神灵庙；THBWiki与夜雀助手均为人间之里/旧地狱/神灵庙。本表采用夜雀助手数据。"),
        ("差异3-西行寺幽幽子持有金", "B站2024攻略写1300-1800；THBWiki与夜雀助手均为1500-2000。本表采用夜雀助手数据1500-2000。"),
        ("差异4-名字写法", "B站旧攻略“二岩瑞藏”按正式译名写作二岩猯藏；“路易兹/露易兹”统一采用夜雀助手写法“露易兹”。"),
        ("差异5-模组数据", "夜雀助手数据中包含 MetaMystia 示例 ResourceEx 包内容（DLC9，如芙兰朵露、八意永琳、神绮、秋姐妹等），属于联机Mod示例而非官方内容，已从本表剔除。"),
        ("免责声明", "游戏版本、活动与平衡可能继续更新；流行·喜爱/厌恶每日随机；随机点单标签会使同一道菜的评价波动；本表用于快速查询与决策，实际以游戏内图鉴和当前版本为准。"),
    ]
    add_sheet(wb, "版本与数据来源", ["项目", "内容"], version_rows, widths=[22, 120])

    out_path = OUT / "东方夜雀食堂_客人攻略查询表_v4.4.0_2026-10-04.xlsx"
    wb.save(out_path)
    print(out_path)
    print(f"特殊/稀客 {len(special)} 名；普通客人 {len(normal)} 名；官方菜谱 {len(foods)}；官方酒水 {len(beverages)}；官方食材 {len(ingredients)}")


if __name__ == "__main__":
    main()

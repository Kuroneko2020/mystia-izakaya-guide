# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
out=[]
# guest table
ws=wb['客人总表']
headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
rows=list(ws.iter_rows(min_row=2, values_only=True))
out.append(f'客人总表行数(不含表头): {len(rows)}')
empty_early=[]; empty_late=[]; no_addon=[]; guest_money={}
for r in rows:
    d=dict(zip(headers,r))
    name=d['客人名字']; guest_money[name]=d['持有金']
    if not d['前期推荐菜谱'] or d['前期推荐菜谱']=='无': empty_early.append(name)
    if not d['后期推荐菜谱'] or d['后期推荐菜谱']=='无': empty_late.append(name)
    if not d['加料推荐'] or '暂无' in str(d['加料推荐']): no_addon.append(name)
    if '铃仙' in str(d['加料推荐']): out.append(f'警告: {name} 加料包含铃仙')
out.append('无前期推荐: '+','.join(empty_early))
out.append('无后期推荐: '+','.join(empty_late))
out.append('无加料候选: '+','.join(no_addon))
# detail table checks
ws=wb['推荐菜谱明细']
dh=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
drows=list(ws.iter_rows(min_row=2, values_only=True))
out.append(f'推荐菜谱明细行数: {len(drows)}')
bad=[]
for r in drows:
    d=dict(zip(dh,r))
    if d['所属版本']=='DLC9': bad.append((d['客人名字'],d['菜品'],'DLC9'))
    if int(d['游戏积分']) < 0: bad.append((d['客人名字'],d['菜品'],'负积分'))
out.append('异常明细: '+repr(bad[:20])+(' ...' if len(bad)>20 else ''))
# unique guests in detail vs guest table
detail_guests=set(d['客人名字'] for d in (dict(zip(dh,r)) for r in drows))
guest_set=set(r[0] for r in rows)
out.append(f'明细覆盖客人数: {len(detail_guests)} / 总客人数 {len(guest_set)}')
out.append('明细缺少客人: '+','.join(sorted(guest_set-detail_guests)))
# B站对照 counts
for sheet in ['B站2024攻略对照','B站2021攻略对照']:
    ws=wb[sheet]
    out.append(f'{sheet}: {ws.max_row-1} 行')
# version sheet
ws=wb['版本与数据来源']
out.append('版本表行数: '+str(ws.max_row-1))
(D/'final_audit.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

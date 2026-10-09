# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
ws=wb['客人总表']; h=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
names={'稗田阿求','蓬莱山辉夜','星熊勇仪','蕾米莉亚','立空汐'}
out=[]
for r in ws.iter_rows(min_row=2, values_only=True):
    d=dict(zip(h,r))
    if d['客人名字'] in names:
        out.append('==== '+d['客人名字'])
        out.append('加料推荐: '+str(d['加料推荐']))
Path(r'待整理/2026-10-04--东方夜雀食堂攻略表/中间数据/check_addons.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

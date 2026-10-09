# -*- coding: utf-8 -*-
"""把最终 Excel 转换成离线可用的单文件攻略查询网站。

- 不依赖服务器和外部 CDN，双击 HTML 即可使用。
- 数据来自 `成品/` 下最新的攻略工作簿；如需更新，先重跑 build_excel.py。
- 输出：guide_template_wafu.html -> 东方夜雀食堂攻略查询_和风版.html
  若存在 脚本/wafu_assets/ + asset_map.json（由 build_assets.py 生成），
  会以 base64 内嵌客人头像与物品图标。
- 琥珀夜色版（无美术资源）已于 2026-10-05 应用户要求停更删除，
  其模板存档于 成品/旧版/guide_template_琥珀夜色版_2026-10-05.html。
     和风版若存在 脚本/wafu_font_subset.woff2（由 subset_wafu_font.py 生成），
     会以 base64 内嵌霞鹜文楷标题字体；缺失时自动回退系统楷体，不影响生成。
"""
import base64
import json
from pathlib import Path

from openpyxl import load_workbook
from pypinyin import Style, pinyin

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "成品"
FONT_SUBSET = Path(__file__).resolve().parent / "wafu_font_subset.woff2"

JOBS = [
    ("guide_template_wafu.html", "东方夜雀食堂攻略查询_和风版.html"),
]

SHEET_KEYS = {
    "客人总表": "guests",
    "推荐菜谱明细": "details",
    "菜谱速查": "recipes",
    "食材速查": "ingredients",
    "酒水速查": "beverages",
    "普通客人": "normal",
    "4.4.0新增内容": "new440",
    "B站2024攻略对照": "guide2024",
    "B站2021攻略对照": "guide2021",
    "版本与数据来源": "sources",
}


def read_sheets(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    data = {}
    for sheet_name, key in SHEET_KEYS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            data[key] = {"headers": [], "rows": []}
            continue
        headers = ["" if cell is None else str(cell) for cell in rows[0]]
        body = []
        for row in rows[1:]:
            body.append(["" if cell is None else str(cell) for cell in row])
        data[key] = {"headers": headers, "rows": body}
    wb.close()
    return data


def latest_workbook() -> Path:
    candidates = [
        p for p in OUT.glob("*.xlsx")
        if not p.name.startswith("~$") and p.stat().st_size > 0
    ]
    candidates.sort(key=lambda p: p.stat().st_mtime)
    if not candidates:
        raise FileNotFoundError("成品目录下没有可用的 xlsx，请先运行 build_excel.py")
    return candidates[-1]


def build_char_pinyin(sheets):
    """收集数据里出现过的所有汉字，构建“汉字 -> 多音字读音（无声调）”映射。

    运行时由 HTML 里的 JS 拼出全拼、首字母和模糊音变体，避免内嵌完整拼音词典。
    """
    chars = set()
    for sheet in sheets.values():
        for row in sheet["rows"]:
            for cell in row:
                for ch in str(cell):
                    if "\u3400" <= ch <= "\u9fff":
                        chars.add(ch)
    result = {}
    for ch in sorted(chars):
        readings = []
        groups = pinyin(ch, style=Style.NORMAL, heteronym=True, errors="ignore")
        for group in groups:
            for reading in group:
                reading = str(reading).strip()
                if reading and reading not in readings:
                    readings.append(reading)
        if readings:
            result[ch] = readings
    return result


def wafu_font_css() -> str:
    if not FONT_SUBSET.exists():
        return "/* 未找到 wafu_font_subset.woff2，标题使用系统楷体回退（可运行 subset_wafu_font.py 生成） */"
    b64 = base64.b64encode(FONT_SUBSET.read_bytes()).decode("ascii")
    return (
        "@font-face{font-family:'WenKai Subset';"
        "src:url(data:font/woff2;base64," + b64 + ") format('woff2');"
        "font-weight:400;font-style:normal;font-display:swap}"
    )


def brand_seal_uri() -> str:
    return "data:image/webp;base64,UklGRuYaAABXRUJQVlA4INoaAABwUwCdASqAAIAAPkkaikOioaEY/R60KASEtQBiEMzrfyZ/BflL7LNf/wW8n1x5i/QHnd/2XqP/SHsB85D9yfUZ+1/7Ye8n6OP776g/92/4fWe+gr5c/slf3b/v/uN7XnqAf//MnuDP1z8gPNvwu+Z/2b9qP7x7emMvqX/pv1J9hf5P9j/w39v/cr/Be03+u8EfiN/N+oF+Rfyn/Pfln7uH0fY8aR/pf+H6gvsH9K/1P5f/5r0fP6z0E+t/+x9wD+Sf0//T/nF8Yf6nwb/s3+d/Y34AP5Z/W/+F/jv3e/3P0u/yf/Y/z3+m/8P+y9rP5//f/+h/kf9F/3P3//Aj+W/0n/Xf4P/Lf+H/If///x/eT7Jf3L9kn9cv++7Zc5z4/vakzQz5giVYlyjBzxPdR6Q3kdq23gWz0VrdK0aO0lVJ6gdXjFgPoo1cN7mY31/daj9s3b/+FZJ6/pAKg42mfBA7LJP22nnW8br4Ow/FuBqP5XmMPE9eBJRHR5v6xxLMpXKxiM2wuBt8qOXve+cnTm9/09ETrXGczyMXXVsOlkIa5nORNBibmjhJmNq+72G3Kb6lmFUvwHoyN0UdU4HzVTF5/ml2G+6qfleV9kOGy6S6z9rCTI5IzZos4Rb3JUdO/slnxo1zFfjt078r55blKixALWnlGVXMxdSwwb+Vz05CBFqlTRjB6H94nP3RZJtM/SOi1oTb5RtmfHrRsrsWwGbpuCg1HWUgpVcthOMnZ2/fkFmGb7C0u/tbfpqs1495L54kVNW2eSZmNgy/3+zEmfucTZExuYAqDuM5vTzzl0CIBKJwoj/qQK3/25Ya+p7P7NLohkYWEmu4ps4nRNTAiE6dsebQBj4dDL40SVlfXg32z0q2ErMg3tl9EzSFKbYjjnaD4ToAAP7+nGSx/49mAtpPV/zxvXel8xWs9Jusp/nvzudtzCBEkddyXKuVd2LZlUu9SGwTIGT7TKmBL7MuYa5AgKzNNPMWZCj8TlPZ1BhwGFhNWsVttiqAb8Lr61t+oxMeiLYyQfdJYSJv65gmSieKz03zD6HKx7oNSuVgfjbJHjq59TcCknR1woYdy/X3Y5r2RSCoqenBjdzGS2t2yfcovcUHIQrOz98J6OvNzJv8AoFGJ3aALBeCU7oHcR96P/mkHYgttE/+2YAIGM+Mp3kcy8xe5Yb8WA/O0hR4FjVJKo5R28hcKvnbVd0R0FvMDlgyaTFMO4RPIYQ+vRB4xqcmcN8R7zo4t4bZ5A+ngIG9sKJFvorw9j+l17DRqXRFVmBB2wZNsPG4ixqXkvK8+QDKGX/icPI4/JX5r/ICNKxbIetFMovJHDZz4j9WL/Qk+YKqzlqiqLRdGSRISOWFe4WiO5mLQB+FhMogit1Pr2IxybHuVQFoWLDmTBqgiOxZS/7QbH8ov7yCav76COn9CiPM0eBCLk3XBgj8LiMQeWjg9ImovDfJZj0R/T7FwvTKZ6SkX5M7QDFr3pPod/jOlmSG1DWTPHUC6CjuezOW0H7qp8NzIrxtdNBUE8o2LXESV5nrjw3v73zMNtzf8NaamIhATKOF6sVHoj3DJg3om+NTI0O5gq56m36jGGo2AcHh2+3+BG9dsLdAR913gvf75JIUgc/dHjnw1dSg4RKT0p2N1M9N58Xd6u289fkrnEJyAJ4n64DKKZa1NuvOXuHqGfHdoV5LRiMifIlV4WjEaUdlt81WEVq/ZbTYbslumQQfmmge1MjZ6fqdYqaNRqNwCXtZyTwBksAJrfCHtzCDQiT3xanUOeSv3tgSEmioGcxPYnSiKhodiBBwfpieeYjG6sI6nvfdm4ALo8doYfP32FZWqjHGzBewbDuRLEFQjGjPIopfw0+2FIEORu+vhTRvQIR5YYuK+MyZ9MCi7C/dFeRr5C2gCNWLcpdmjq5IzcvkZ0EwRoOTlH3btFOjFqTPDKq1U3uBSAx7sCNqPoU4+3nAHk4O/5J3S+YYyLSXYAaCCnc/Y/8ag1sCizrMSCVH2VwxDAZxEvtc0HS70CPvJI0AZSvf1IRbnkR5gcwyTQ69z3op8wdC4CbnfoFFJZWT6diRhWbDaeye1WuQYTIu50YwZkEh3prBXLYLTPXyIz5DrmhaZh7BYZAK8CSfQIsR/qXv3CHuyS9HYZykMs/jz3qZBK6Ao1WfonBqFWCyE4CUsBnzCF2Ks3n5FFSpkYlKYi+ZDkxE9qVDZBMJ/Ju41DPpsMOUV9gYnIaVN5ZfLDKQXZKJ+kzTaD5MtVxMcNWLtHlaASMkyitL76xl+8oLEydHvEcyGzMDA0vHbccb0UulvAzSKN3ccJ/81t+fsVqwGbIAx9i40PAB0hnj3F6bHJyjS2z5F5b4PFg8FBzMCMoszu0UDqTKqVMAHTwqOjU/QqL5TObHZXj7447nfUQ1r/NHnJP2kX8s1xMOC3PIOj86OUae7dVzgMJF4KV41c/FB0hQctf9QqF6RI7zfuEOOpBJS9H5F36DX2lkiitn5GO5mJUx7PXJBPKxBXFsifynViKEZEOfu4KbDeawr0sDN3M+j1b45zIlYds5DVtClathANcF/fFnXVX133fbQ4ucvHARe5cLqFO17vlz43lNevr/XOv6zxbWRhkdlHpHSLULby0MO2g/O3EYpAMnIOB8OdYVV2axs2glm4Z+ZX61tmY4OQCcZDIXz4+o4sOIRc8+dVQc2Mi39DmwFy6Z9ClI0zi3PMIpBeOII+bJYl597ZabHcTWog1CJHCy7D3zHuyOSMzUarMExE5o6V2piQ5JpE6JtEQGTFUj5TmwR457+a2A1toKmjBXnJDaiuz7Bh8xmWTvuHYeumKMDkgDVERigsuUveX9hsMFWMvkH4WLi/+uUC+WFO4B/2n7TTvGnntsOyFR18hJa/kVUN67fOBocZzNRoi3U8eVFBvE+lXO+kS5pPVfp57ImukYKSjFhGaDxoojEeTut8n3M2XLxxehRqxJ0baHmY1adcCVUAFvvRjGzB5poVIVyUaqcvS1aCTnjQv6Fxb43nsi390qkaxdsQqALjjtHmvStQleOAM1iepxmL5OmYw/pbQGn0Vs/Ocq4eLfkZNsVHG8yvXjqbMXilkxSQeTlRTBlXaBDYnujARairlqvNN0CbBuwzJ3RLxpas9PckyGNswxry4SIdzCbgi9McYbZ/HPhQIDaUtxvqa1i4DSKerpkzmRcr1DK+4VuxWdzEFfbLEMmJtUwSPT9bQVoSIV1sjjyZwz7EFb2lcmXuWPohb7HoCOxaqQvLHSfQ+ITIpsDdrGbdBWwKHHgDVAIc+lhEkQqJIEgDh+/AkLfGpbGvdrL8qqPy9geJl+9dSDJDwrMY0BjX1lTyFr4333Kix+Ov2FVJ75hbUUp5M3kWuWe3Pkagg7pc2x9qqrfUpQOfox0GF+/7zBz6XmtfffpnMzkhs8tKScS02rfzDT+Zvlqr3yqR9Hbe5CiCVAIi93yeb89YXQwWcU+5OQthtyG/wRGHFs3drCfiaeyURsn+u8mRgkqQsz4t/avqUhQPw2io2K/ugE7pqP/WS/YvdXpTUtbLHc4x1RGbTSnO3EhdT8p1snrWC1gE6TXCgujHFy1eRQvkrqahNpGJ7IaDdYNYjnxX73SzhPxwEBtUHkoADiI7MHldCqJ61WH1usrn1rVxC2EOlNJKhc2Kv8aPTxAtz7OseZI4TFAguCzpl5vnzzaCuQZODKqR+PEBRt1d1pBLNWBroikMQfniJTiVA2P5RmUaz/HfDU3k/tLm2tlKm/eN7cBFG8SN9znOcK3Qhm7VGaDV0kHjJKiQHH3tyKAjZVy3Eyl9n2/8/Mr7tzZmkM2b9pXqYZ+t4KPySZOoXxT3uXymXxJ29YBIsGW4mWJ2kT8lT8r6IGxdWKTrUdnuoPaA7jGRkS7lckpb5hNgi5j7tGBFCua6DZ30Y0896HhO51xHn37lhZD+45d0exS7qbYWm9A+kYyvmGyLhfp8iw3J0vbs6mXP/KUQSwlAchjwhQZdnZZtpWU9L8+5u2vR2pn2CdNIdHuXA3imNHEreacgElhOKg3gC0vXUR34mVxHrEEZBCv6uswnhAL0b0I/Fgul8LcYHLnd+SmRhBsCAqPAHzIbxpAlblbpDYOXEVH07FRuefK1Cj24bzWWDQE3+HHwlu2JDeD7dBpM2KM4COIJOJxscm4qlJX8zO9Sw7a+CBm6ZbGxhS9kXf7qjyd5ep9b6r7StnvDU8LWzg8ioI4dV1scfJu3L+ScVtWAJkMAE0E635osPXUaK0CG1a7Hns6E3JtIpozsLqv2QPL7FA9ILMqRmjnF5FzxCGFoL/FUz5dp1mPXbTDQznR32JfIth6HrSjC0WdlAxq35zmKe/N6qQAZRul+dkMxW6rh4N/3EhX1qSV0gF0xxwpsEvRmfCXrCwSqSo5YjaIp812A67mMVhmoduJv4ymWKpAtkJXpjVWbnolXYbqOoNc6/UlJ/qSjO+1OX3s/DIVL8IYehK1w0//wql41VjGT9l6ksuC0kwmbQSAvz0AMt3McdzAs/e1hXp1S+fZyKiXgovFJrmqNVEdMJXgtGQyGMWCBxWke8zJZG3crnDM+nhE+lByizzazH/TRwG55/Fkr9zGFgfpzfqVST6y5uSZAolXC5mHSW0Ei/DP4UYaFPR5kYOsDWNfG/PyoTp1Cxv2NIlPN9FwfHRAbSUxahg4qR1jwrpNqSM5LIg20/9Cu+CdgWKBH/PYK4bbP5lanY2gLlPF8T6BD0kcq+CKfJzo0/Lvsz1py368qIFQXqJjB6R1YwuOxUqYh3UWjyYsV50ND8kZZ/0+mQ25vOIPZfm8Qb0l/7mYMxS8FAS8sN/v1+hG42mLcOcN8PXjwtdbu7ESarjdLo0/iwvt6u4B+eSlUoIfmEu1ma/3SI/DpKRhdBzlvDURkvRTuIEOWusWh3QEtuVV7zW4+AAKwMHdSHgfS50XKzMNMaGQnRsGlVHQPmN4XWop45dsngPepMF9gnZjqAtuibuDPNPaEhet7xs5Mp+gAkAZeyEOkemU6Aj8g7eopMqUWQ8Abe6KWtup2MPJMw5yhPmamNsZtl/HDjIR1jejP8qqb/u2QIT9E8JMWWYloWp4lUZsW6COdu/j88c+YcTrBmNL5le/GVr/MRAzLSjhiS8gN1HDlko+jNYV711ucldbjqNaRzJvVGBV/1uiz/Q2F3lDQ5FiboBbq6q7NcAtzNsFF8fg222M536qr8C+PTAgnUf80e7caMzU1f56IH+aKLnXJxyC2uD80K4vpKGV121R1Rysyt7cMmIuU094RwVeomM8YyA3V8eWUpTdNU8vw+xOn5Ih4zp3t8rOd/CNVWdm/yyjw1gWBLKZvtTLqof85qf7Hjj/2iDrwAXJ57h46hyqRwLIL6PS+eBqvWPOmOFt1pbTRZG7FuBI8U3TbTipiTI40rctbqocMT3bSCyLtpcHIEnoQd+7cSAJqkOB4c9numnXCDDYadbjj4noJwj/ArXWBp9Q0xlcejN52qHGvbFcWI0tiaFGC7/TgmQdS29LZMUUOtVUYLuKvm8VDvpCbg3USo3PJ6uzxYiwJ7qIGnybqHpm1hSpNXVDz5AukQ2DNl4enegKMLAn13zoTh1wgdyEUHESJH8SG9GwOrHukLAJtUmGilt1qcuJ0nMKzOkRbpJ0inx9uBQWvTAoWD7G59MXQdKQoYWW1FdMVea7kLkJwPZ5Bw9dyFjaWKHMdqYSOY2acQPbEV5iJrwCeKUluuU+7XGFp+8FEwZurotiF2BzTbcIgJtATFJR22FPRgoR62QuWflF6h+O1eJc2ndHMWTJVPbEorc1I9lQA1TogUrZCI66wd2Jha8/0utHF53w3exRNOvXAAHh/G4+wuCuTvvbKSgn/uHBgBh+rVSgL7hUq3PNduJ2frvgy/B9/K5rotmjN5ZUiKjmKEewQUPRLnwkBOTR/l9xhDwN0adOrAvZ7WE55HNnX5dZ9Bbz2l6/a4AZljjpu1DzSVrFq+h60xmAveFBoBJPCwiTl9pb90rQgov5kwu69OtmTlCrNeY4IRlzGwTI5TQOb+TE17/P/QRxBpwoBN32U8oY1G/WRk8OWXrCR3an+F7T/7QDEmnw6MzKebK3B2+Oty2/vnXxGmdp0FGhgiD9EI0B6wJBLid9hwtJaH0LrgFx6OxQjnqdm+DadjjqFB75Z4B8AraixPr8RaQUA/5aZp5IhHdd0REEvk6F9Fw8UExoYUOOJZwt0/CKhCwT39jCXOvITlIj0sbnUI17XdD6vhZk8IGrDXTr8dFy+NKb0eI8O547OSUS5bPxj6WsuVzdNd9cNpJNuc6u1TnTXLmEBNj78iy5dwyoucmUAjSC5h7YoNT3wH8XigPgaj8+4eeAFQAyaHRqs72X4y6tKGC8MgCTy4mwjWIkXE9+TYkDL+vmcjVG9kIB1Uzlvr2BIB0iCx8g1p4SIanlvISRwFR57ONgeGoXoRtB84nmL1o5fyYzxvo+OM3eeqZ7Vviwf3JaCIWNxXUyQof2ZcH6CDLfPSoFL2hQ6TQgoQu6A+1zSUufhgPlYraHfnAcR72yZuIMu6YFxynqUxuEcUtuR9zoiVM16zcpc9qBadikHcJSrp0mzPUnwz9MM9a+kkvpvZcyFZShnflPA/R06L1yfS06q4OXaJ301AACN9HvONqeC2czGbWgQT+P6R5lelI7juPSray2Fcxbw6aod8FMkT90EQt4+JhL7l+uuqr8mHjnjn2xqBfgSkqzxX1sm82lRqderek8/3CAc7MlCYY9rko//kMfn/VJDTz34Vd9FU1pTmdh07QfzF0jzB0UBPC6VDE491/HozdzIUL2lKroe8SEa6+SUzNNtmIoS/G3DKaKG0kiZIFhzVbjO4o9nO2Bz+9kfeUFfblOJMyHckBMwQqYoxL/Jy6YYmt9dXZgIOiLVMfoRsFIHjsROLcda57hCectec+R/PPdavAfM+yCkm32kCs2Z0dW/YOJjPehw6t/DqTjhF+z0gqBpXos/bj4us+Zlvrln+XUmDgtRaEP3PydDGeSDKlPCYW3LEMDj+tEryOtre1uBwTXEJeZi31H/58LMHa7zyN7CuluLbXYhw2H4+Zu4rtzBuc6LYW/C/8YC9Z6yBAHrvfbGbPf3ZwyjdeHVcOPUh7fqbdRFgoClATkrFu2fa9pI23mylGf4jdAnYooUY7V6QyWiZ4SaIsyV2Py9BuP+hT4jKp2swJyfYDugCYSIaH+PUVskQvbYZjBXlAB5mdYhuXIIOMvnjTT3pAeNFXthckZitXmbrEnJrj1HEQhxaV+skx2QtJQcTyZBdg9ZPdiESZT0v98bw90K4h5UlR76JJ/TW1blh6wdgRRVwvbpJBiCaUHwGO7uSQfxHflCEJM00yqOWEU4/Kr/K+gTdSxLZvD5PfHcaNfuu5+fX+uxzyLJwNOkkGofZCcJDlKiz8KRtH3FbT7bKTohn5ogw5Ff8k6Pe/uWBqOxKkkYpQeAeXKdgVmZMZaEqXDYo7okAghVAdMihsOdfEIgvixPhW0W0zt77VlVkyB5JDxR7VbLzAOEq3DxfUtrmD3guzImreX9rQdlTzRyZBFJC50dGUUT4eABtLwAIx/SD6iNR5Ymv8Ju9koF+EZ/RSHN5zD+5iuEGtJNZ6s4xM+EaVDBMEdwZ54VwLWdgPGwawGo6BOgCd0jBN/VlBKBd2Zg//Xlbo3pjhzsYporedwnZG3y1oEGmbuIFvd3RgjTsry+x0iFEjD6QW1X0h9xNW514PC0d+jfWp1Kr6YSThIEEvQ8/UepmN/h4ygNqZjRTZvMWvvxdpcMP5hmWFf90nsF2ZJdyfdzd1g5g6yVSAmL/hShCFD2caCEKN6KKLsQ99qpVUF76z3WdwT4KgXNloHoIObYL3MSCTOfGDidYbe8JMb+FOPqZSS2GGUBK0RryUGUHDRUbVWD5et0KWcrlktPXn9LyPI3B4A3AtCUKDqqZTWi8bcdAym9+VC48B5jOHPStyLWxRNnU4TX6o+Lum0hCTEoEJKb0BVzCOqQbnkq2wcKvO3BL6Bn/kyMdr/mKS7vuwmBtqLL7w51uyeLtjXyWjUxdt8d+9AYlXykmVAypaeODJi67yzTJAeU6BV1djK3CH8GmJ8IIukKEIqagOpHcsR7xthFPqhcHXvgMYHX4xfduZw0QRjr+Bvc5KMEO13J9Mx0K5NNSZHEC+sqKYTv9Nj8MJn1pbuHlFRXyWe4SWlvwaj6zQWPzRvxFwAsxaKPAGIyjJVeDkWcR/rb9N451GsQm+2lkcKMQFE4u+fS6xYOEl0Uw7P8RjCwqDnfyw2OqHI1syDB1NjpQTOWhm1LAm37dksB8pZFqBN4hs/jDYD/+3tCL+Q30paUBda8sVBVQpPxf1+6RYiCEIY9EGGX6wYab0NeVoANrlbblznLLwhq4oNuEMPE/5NxJ+rCSRyGdyETjBpPMWpB34hSOYgw0awj50vzwd+bLNRRxsS1HhQzS/mDhbB4Ro7ia69Pp1dfabI9fJVIS7J59LAWuyH51kUB3++fWJqBpJiSbynZmWyPlj6v38zJXluKOgeUq2uXYC0RX9LMTxEakT4Vmzuaz0Js4wU8ym4m3UVLMHUGPnuda/bfNpe+8lVfGVHKgbshtaT2/b5nPvqdROtimLggAGpzdoauaE6Qa0aub3UlMqZKOO79N1yeg4C5Iz8PcLGDpWpy3FWfNvgXUVYeQF55UGy8e5zDSCmyTlQ2nX6TgUkr0vAMOGUyJVDbgBQQDCA8Mu0EEVatMK1ZkTDY7iblTmdj+5JoMReYPHTBRCMylfDBFuCsCjV7JAzVnl9mYb/mbNycgiiJYJcWYceb66kUBmCFHEE9b1uzg9eSPnocKQK6LhU7zXoPa3FeS+p14HWKW41XCouFJbnyasSmrmtGViRlABHn/D4cWvctEf89QvekyU6fKVie+x65qCOF48E7VN2Ae5lbrvUPRGUoeea0U5TFMetiLp6Xgu8Pw43HQNo47IKDoT34Vv4F6E0dacAoU8hTJMDoQ2Xt2GcB/Ad6ooX0Jn2GgwKwkUOKTajMJt9TdBR/WORETs6AJE4DxwDzJusZEfzzCFOZojQyWtVbWHf4JSS3DMAt7mkAYD26as1634AOb6Tjcpw+gHXinH10RrGcNBsReMz8k3e5RLYAAA"


def wafu_assets_js() -> str:
    """把 asset_map.json 描述的头像/图标全部内嵌为 data URI。"""
    mapfile = Path(__file__).resolve().parent / "asset_map.json"
    assets = {"avatars": {}, "icons": {}}
    if mapfile.exists():
        mapping = json.loads(mapfile.read_text(encoding="utf-8"))
        for kind, names in mapping.items():
            for name, fname in names.items():
                f = Path(__file__).resolve().parent / "wafu_assets" / kind / fname
                if f.exists():
                    b64 = base64.b64encode(f.read_bytes()).decode("ascii")
                    assets[kind][name] = "data:image/webp;base64," + b64
    return json.dumps(assets, separators=(",", ":"), ensure_ascii=False)


def main() -> None:
    xlsx = latest_workbook()
    sheets = read_sheets(xlsx)
    payload = {
        "meta": {
            "generatedAt": "2026-10-04",
            "version": "v4.4.0",
            "excelName": xlsx.name,
            "sheetNames": list(sheets.keys()),
        },
        "charPinyin": build_char_pinyin(sheets),
        "sheets": sheets,
    }
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    # 防止文本中恰好出现 </script> 截断内嵌 JSON。
    raw = raw.replace("</", "<\\/")
    for template_name, out_name in JOBS:
        template = (Path(__file__).resolve().parent / template_name).read_text(encoding="utf-8")
        html = (
            template.replace("__APP_DATA__", raw)
            .replace("__GENERATED_AT__", payload["meta"]["generatedAt"])
            .replace("__VERSION__", payload["meta"]["version"])
            .replace("__EXCEL_NAME__", payload["meta"]["excelName"])
            .replace("__BRAND_SEAL__", brand_seal_uri())
        )
        if template_name == "guide_template_wafu.html":
            html = html.replace("/*__WAFU_FONT__*/", wafu_font_css())
            html = html.replace("__WAFU_ASSETS__", wafu_assets_js())
        out = OUT / out_name
        out.write_text(html, encoding="utf-8")
        print(f"{out}  {out.stat().st_size / 1024:.1f} KB")
    print(f"数据表 {len(sheets)} 张，源工作簿：{xlsx.name}")


if __name__ == "__main__":
    main()

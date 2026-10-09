# -*- coding: utf-8 -*-
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
ws=wb['推荐菜谱明细']
h=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
c=Counter(); per={}
for r in ws.iter_rows(min_row=2, values_only=True):
    d=dict(zip(h,r)); key=(d['客人名字'],d['阶段']); per[key]=per.get(key,0)+1
for k,v in sorted(per.items()):
    if v<2: print(k,v)
print('distribution',Counter(per.values()))
D.joinpath('中间数据','rec_counts.txt').write_text('\n'.join(f'{k}:{v}' for k,v in sorted(per.items())),encoding='utf-8')

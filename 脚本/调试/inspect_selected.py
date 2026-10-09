# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
ws=wb['客人总表']
headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
want={'秦心','琪露诺','博丽灵梦','西行寺幽幽子','饕餮尤魔','八云紫','蕾米莉亚','火焰猫燐','茨木华扇'}
out=[]
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] in want:
        out.append('='*30+' '+str(row[0]))
        for h,v in zip(headers,row):
            out.append(f'{h}: {v}')
(D/'selected_guest_report.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

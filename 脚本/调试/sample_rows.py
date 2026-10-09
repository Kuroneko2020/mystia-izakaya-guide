# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
out=[]
for sheet, names in [('客人总表',{'西行寺幽幽子','绵月依姬','秦心','露米娅','琪露诺'})]:
    ws=wb[sheet]; h=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] in names:
            out.append('==== '+str(r[0]))
            for k,v in zip(h,r):
                if k in ('出没地点','持有金','喜好料理标签','厌恶料理标签','喜好酒水标签','前期推荐菜谱','后期推荐菜谱','加料推荐','推荐酒水'):
                    out.append(f'{k}: {v}')
(D/'sample_rows.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

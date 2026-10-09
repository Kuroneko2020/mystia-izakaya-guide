# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
out=[]
for ws in wb.worksheets:
    mx=0; cell=''
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                l=len(str(c.value))
                if l>mx: mx=l; cell=c.coordinate
    out.append(f'{ws.title}: max_len={mx} at {cell}')
(D/'length_report.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

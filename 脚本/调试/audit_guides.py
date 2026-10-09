# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
out=[]
for sheet in ['B站2024攻略对照','B站2021攻略对照']:
    ws=wb[sheet]; h=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
    rows=[dict(zip(h,r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    out.append(f'==== {sheet} rows={len(rows)}')
    empties={}
    for k in h:
        empties[k]=sum(1 for r in rows if not r.get(k))
    out.append(repr(empties))
(D/'guide_audit.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')

# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
ws=wb['客人总表']
headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
out=[]
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] in ('茨木华扇','博丽灵梦'):
        out.append('==== '+str(row[0]))
        for h,v in zip(headers,row):
            if h in ('喜好料理标签','厌恶料理标签','前期推荐菜谱','后期推荐菜谱','加料推荐'):
                out.append(f'{h}: {v}')
ws2=wb['菜谱速查']
for row in ws2.iter_rows(min_row=2, values_only=True):
    if row[1] in ('饭团','蜜汁叉烧','无意识妖怪慕斯'):
        out.append(f'菜谱 {row[1]}: 标签={row[7]} 价格={row[4]} 版本={row[2]}')
Path(r'待整理/2026-10-04--东方夜雀食堂攻略表/中间数据/debug_guest.txt').write_text('\n'.join(out),encoding='utf-8')
print(f)

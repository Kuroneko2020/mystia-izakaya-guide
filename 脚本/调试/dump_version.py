# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[1]/'成品'
f=sorted(D.glob('*.xlsx'))[-1]
wb=load_workbook(f, read_only=True, data_only=True)
out=[]
for ws in wb.worksheets:
    out.append(f'===== {ws.title} ({ws.max_row-1} 行数据)')
    if ws.title in ('版本与数据来源','4.4.0新增内容'):
        for row in ws.iter_rows(values_only=True):
            out.append(' | '.join('' if v is None else str(v) for v in row))
    else:
        out.append('列：'+' | '.join(str(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))))
(D/'sheet_summary.txt').write_text('\n'.join(out),encoding='utf-8')
print('done')
